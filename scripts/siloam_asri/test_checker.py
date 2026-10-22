"""Uji checker: rusak salinan file output, setiap kerusakan harus tertangkap."""
import sys, os, shutil, hashlib, openpyxl, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import check
from openpyxl.styles import Font, Alignment
OUT, BAK, TPL, RAW, CORP, XL = sys.argv[1:7]
S = 'Siloam Asri'
def cellpos(ws, uraian, kelas, k, j):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value == uraian and ws.cell(r, 4).value == kelas: return ws.cell(r, 6 + 3 * (k - 1) + j)
def setv(u, kls, k, j, v):
    def f(wb): cellpos(wb[S], u, kls, k, j).value = v
    return f
def flagcol(ws): return ws.max_column
def mut_flag(u, kls, fn):
    def f(wb):
        ws = wb[S]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 2).value == u and ws.cell(r, 4).value == kls:
                c = ws.cell(r, ws.max_column); c.value = fn(c.value)
    return f
def other(fn):
    def f(wb): fn(wb['Siloam Makassar'])
    return f
M = {
 'tarif lab': setv('CRP Kuantitatif', 'III', 1, 2, 524001),
 'tarif kamar': setv('Ruang Rawat', 'II', 1, 2, 300000),
 'tarif visite': setv('Konsultasi Spesialis Rawat Inap', 'I & R. KHUSUS', 1, 2, 160000),
 'tarif IGD': setv('Konsultasi Dokter Umum IGD', 'III', 1, 2, 280000),
 'tarif pompa': setv('Syringe Pump / hari', 'II', 1, 2, 221000),
 'tarif admin': setv('Pasang IV Line Access', 'III', 3, 2, 309000),
 'label VIP': setv('Ruang Rawat', 'I & R. KHUSUS', 1, 1, 'VIP'),
 'label Excel lain (OPD)': setv('Injeksi IM / IV', 'III', 3, 1, 'OPD'),
 'label ED -> UMUM': setv('Konsultasi Dokter Umum IGD', 'II', 1, 1, 'UMUM/NON-KELAS'),
 'label kelas tertukar': setv('GDS', 'III', 1, 1, 'STANDARD'),
 'nama typo': setv('GDS', 'III', 1, 0, 'GLUCOSE AD RANDOMS'),
 'salah ketik dirapikan': setv('Pasang IV Line Access', 'II', 4, 0, 'VENA SECTIO'),
 'spasi ganda dihapus': setv('O2 / Hari', 'III', 1, 0, 'O2  PER DAY'),
 'akhiran kembar dibuang': setv('Konsultasi Dokter Umum IGD', 'III', 1, 0, 'ED CONSULTATION RMO - DAY'),
 'akhiran pada nama = kolom B': setv('Ureum', 'III', 1, 0, 'UREUM - H. LABORATORY'),
 'akhiran tabel pada nama unik': setv('CRP Kuantitatif', 'II', 1, 0, 'CRP QUANTITATIVE - H. LABORATORY'),
 'akhiran salah tabel': setv('Konsultasi Dokter Umum IGD', 'III', 1, 0, 'ED CONSULTATION RMO - DAY - B. ROOM RATES'),
 'kutipan palsu': mut_flag('Nebulizer', 'III', lambda t: t.replace("'NEBULIZER TREATMENT'", "'NEBULIZER THERAPY'", 1)),
 'kutipan teks Excel': mut_flag('GDS', 'II', lambda t: t + f" | {len(t.split(' | ')) + 1}. 'Glucose Ad Random'"),
 'penomoran': mut_flag('Thorax PA', 'III', lambda t: t.replace('2. Pemetaan', '3. Pemetaan', 1)),
 'flag beda antar kelas': mut_flag('Thorax PA', 'II', lambda t: t + ' x'),
 'flag bukan YA': lambda wb: [setattr(c, 'value', 'TIDAK') for c in [wb[S].cell(5, wb[S].max_column - 1)]],
 'flag tidak merah': lambda wb: [setattr(c, 'font', Font(bold=True)) for c in [wb[S].cell(6, wb[S].max_column - 1)]],
 'alasan tidak rata atas': lambda wb: [setattr(c, 'alignment', Alignment(vertical='bottom')) for c in [wb[S].cell(7, wb[S].max_column)]],
 'urutan triplet beda': lambda wb: [setv('Creatinin', 'II', 1, 0, 'UREUM & CREATININE')(wb), setv('Creatinin', 'II', 2, 0, 'CREATININE')(wb)],
 'urutan triplet beda (triplet utuh ditukar)': lambda wb: [cellpos(wb[S], 'Creatinin', 'II', 1, j).__setattr__('value', v) or None for j, v in enumerate([cellpos(wb[S], 'Creatinin', 'II', 2, jj).value for jj in range(3)] )] and [cellpos(wb[S], 'Creatinin', 'II', 2, j).__setattr__('value', v) for j, v in enumerate(['CREATININE', 'STANDARD', 114000])],
 'triplet "-" tidak lengkap': setv('Jasa Perawatan', 'III', 1, 2, None),
 'tarif jadi teks': setv('Ureum', 'II', 1, 2, '114.000'),
 'A-E diubah': lambda wb: setattr(wb[S]['B2'], 'value', 'Konsultasi Dokter Umum UGD'),
 'freeze diubah': lambda wb: setattr(wb[S], 'freeze_panes', 'A10'),
 'lebar diubah': lambda wb: setattr(wb[S].column_dimensions['F'], 'width', 30),
 'sheet lain diubah': other(lambda ws: setattr(ws['H2'], 'value', 1)),
 'sheet lain dihapus': lambda wb: wb.remove(wb['Siloam Sentosa']),
}
tmp = tempfile.mkdtemp()
fails = 0
for name, fn in M.items():
    p = os.path.join(tmp, 'mut.xlsx')
    wb = openpyxl.load_workbook(OUT); fn(wb); wb.save(p)
    e = check.check(p, BAK, TPL, RAW, CORP, XL)
    print(('TERTANGKAP' if e else 'LOLOS!!'), '-', name, '|', (e[0] if e else ''))
    fails += not e
# template diubah
tp = os.path.join(tmp, 'tpl.xlsx'); shutil.copy(TPL, tp)
wb = openpyxl.load_workbook(tp); wb.active['A1'] = 'x'; wb.save(tp)
e = check.check(OUT, BAK, tp, RAW, CORP, XL); print(('TERTANGKAP' if e else 'LOLOS!!'), '- template diubah |', e[:1]); fails += not e
print('TOTAL lolos:', fails, 'dari', len(M) + 1)
