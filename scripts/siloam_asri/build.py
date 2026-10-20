"""Bangun sheet 'Siloam Asri' di file output. Pemakaian:
python3 build.py TEMPLATE.xlsx items.json OUTPUT.xlsx"""
import sys, json, re, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, __import__('os').path.dirname(__file__))
from spec import SPEC

SHEET = 'Siloam Asri'
SRC_AE = 'Siloam Bekasi Timur'          # A-E identik di semua sheet Siloam
KELAS = {'III': 'BASIC', 'II': 'STANDARD', 'I & R. KHUSUS': 'DELUXE'}
ADMIN = {'IV LINE ADMINISTRATION': 319000}   # tabel A hlm. 3 (format '319,000'), dicek visual

def num(v):
    v = v.replace(' ', '')
    return int(v.replace('.', '')) if re.fullmatch(r'\d{1,3}(\.\d{3})*', v) else None

def main(tpl, items_path, out):
    items = [o for o in json.load(open(items_path)) if o['kind'] == 'item']
    def find(name, page):
        r = [o for o in items if o['name'] == name and o['page'] == page]
        assert len(r) == 1, (name, page, len(r))
        return r[0]
    wb = openpyxl.load_workbook(tpl)
    assert SHEET not in wb.sheetnames
    src = wb[SRC_AE]
    ws = wb.create_sheet(SHEET)
    N = max(len(v['trip']) for v in SPEC.values())
    for row in src.iter_rows(min_row=1, max_row=src.max_row, max_col=5):
        for c in row:
            d = ws.cell(c.row, c.column, c.value)
            d._style = copy.copy(c._style)
    for col in range(1, 6):
        ws.column_dimensions[L(col)].width = src.column_dimensions[L(col)].width
    hdr = []
    for n in range(1, N + 1):
        hdr += [f'nama_layanan_{n}', f'kelas_kamar_{n}', f'tarif_{n}']
    hdr += ['flag_mencurigakan', 'alasan_mencurigakan']
    fT, fF = PatternFill('solid', fgColor='FCE4D6'), PatternFill('solid', fgColor='FFC7CE')
    for i, h in enumerate(hdr):
        c = ws.cell(1, 6 + i, h)
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True)
        c.fill = fT if i < 3 * N else fF
        ws.column_dimensions[L(6 + i)].width = [42, 24, 11][i % 3] if i < 3 * N else (18 if h == 'flag_mencurigakan' else 120)
    FC, AC = 6 + 3 * N, 7 + 3 * N
    seen = set()
    for r in range(2, ws.max_row + 1):
        ura, kh = ws.cell(r, 2).value, ws.cell(r, 4).value
        sp = SPEC[ura]; seen.add(ura)
        for k, (nm, srcname, pg, mode, exname, lab) in enumerate(sp['trip']):
            col = 6 + 3 * k
            if mode == 'none':
                vals = ('-', '-', '-')
            elif mode == 'kelas':
                o = find(srcname, pg); kc = KELAS[kh]
                vals = (nm, kc, num(o['vals'][kc]))
            elif mode == 'ed':
                o = find(srcname, pg); vals = (nm, lab, num(o['vals']['ED']))
            elif mode == 'admin':
                vals = (nm, lab, ADMIN[srcname])
            assert vals[2] is not None, (ura, nm, kh)
            for j, v in enumerate(vals):
                c = ws.cell(r, col + j, v)
                if j == 2 and isinstance(v, int): c.number_format = '#,##0'
        ws.cell(r, FC, 'YA').font = Font(bold=True, color='9C0006')
        txt = ' | '.join(f'{i}. {t}' for i, t in enumerate(sp['flag'], 1))
        c = ws.cell(r, AC, txt); c.alignment = Alignment(vertical='top', wrap_text=True)
    assert seen == set(SPEC), set(SPEC) ^ seen
    ws.freeze_panes = 'F2'
    wb.save(out)
    print('N =', N, 'rows', ws.max_row - 1)

if __name__ == '__main__':
    main(*sys.argv[1:4])
