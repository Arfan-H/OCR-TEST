"""Re-check otomatis sheet 'Siloam Asri' (terpisah dari builder).
python3 check.py OUTPUT.xlsx BACKUP.xlsx TEMPLATE_ASLI.xlsx raw.json corpus.txt EXCEL.xlsx"""
import sys, json, re, hashlib, collections
import openpyxl

SHEET, SRC_AE = 'Siloam Asri', 'Siloam Bekasi Timur'
TEMPLATE_SHA = 'e430ea0efb28cc36aa0c14601e4afdd6b32f827c46fb60424296027278e550c3'
KELAS = {'III': 'BASIC', 'II': 'STANDARD', 'I & R. KHUSUS': 'DELUXE'}
EXTRA_LABELS = {'ED', 'OPD / ED / IPD / RJ / RI'}
BAD_LABEL = re.compile(r'VIP|SUITE|ICU|NICU|PICU|HCU|ISOLA|EXECUTIVE|DELUXE A|PRESIDEN', re.I)
COLS = ['ED', 'OPD', 'OPD ODC DELUXE', 'BASIC', 'STANDARD', 'DELUXE', 'VIP', 'VIP DELUXE', 'VVIP', 'EXECUTIVE SUITE']
N_ = lambda s: re.sub(r'\s+', ' ', str(s or '')).strip().lower()

def num(v):
    v = v.replace(' ', '')
    if re.fullmatch(r'\d{1,3}(\.\d{3})*', v): return int(v.replace('.', ''))
    if re.fullmatch(r'\d{1,3}(,\d{3})*', v): return int(v.replace(',', ''))
    return None

def style_sig(c):
    f, a = c.font, c.alignment
    return (c.value, c.number_format, f.name, f.sz, f.b, f.i, str(f.color.rgb) if f.color is not None else None,
            c.fill.fill_type, str(c.fill.fgColor.rgb), a.horizontal, a.vertical, a.wrap_text)

def load_source(raw_path):
    """Daftar baris sumber: (page, title, name, {col: angka}). Judul = section/judul tabel terakhir di atas baris."""
    d = json.load(open(raw_path))
    rows = []
    for x in d['items']:
        c = x['cells']
        if x['page'] == 3: continue  # tabel A ditangani di bawah
        if x['ncols'] == 12: c = c[1:]
        if len(c) != 11 or not c[0] or c[0] == 'DESCRIPTION': continue
        rows.append((x['page'], x['title'], c[0], {k: num(v) for k, v in zip(COLS, c[1:])}))
    # tabel A hlm. 3 dibaca dari teks halaman (baris digabung pdfplumber karena sel kelompok bergabung)
    for nm, rate in [('SILOAM CARD', 16000), ('ADMINSTRATION FEE', 55000), ('DRUG ADMINISTRATION OBSERVATION', 1061000),
                     ('SRT C19 CONTRIBUTION', 103000), ('SURAT KETERANGAN HASIL COVID-19 TEST', 103000), ('IV LINE ADMINISTRATION', 319000)]:
        rows.append((3, 'A. ADMINISTRATION FEE', nm, {'RATES': rate}))
    titles = collections.defaultdict(list)
    for t in d['texts']:
        titles[t['page']] = t['free']
    return rows, titles

def check(out, backup, template, raw, corpus_path, excel):
    E = []
    err = E.append
    if hashlib.sha256(open(template, 'rb').read()).hexdigest() != TEMPLATE_SHA: err('template berubah (hash)')
    wo, wb = openpyxl.load_workbook(out), openpyxl.load_workbook(backup)
    if wo.sheetnames != wb.sheetnames + [SHEET]: err(f'daftar sheet salah: {wo.sheetnames}')
    for name in wb.sheetnames:
        if name not in wo.sheetnames: continue
        a, b = wo[name], wb[name]
        if a.freeze_panes != b.freeze_panes: err(f'{name}: freeze berubah')
        if (a.max_row, a.max_column) != (b.max_row, b.max_column): err(f'{name}: dimensi berubah')
        for col, dim in b.column_dimensions.items():
            if a.column_dimensions[col].width != dim.width: err(f'{name}: lebar {col} berubah')
        for row in b.iter_rows():
            for c in row:
                if style_sig(a.cell(c.row, c.column)) != style_sig(c): err(f'{name}!{c.coordinate} berubah'); break
    if SHEET not in wo.sheetnames: return E
    ws, src = wo[SHEET], wb[SRC_AE]
    # A-E identik template
    if ws.max_row != src.max_row: err('jumlah baris beda dengan template')
    for row in src.iter_rows(max_col=5):
        for c in row:
            if style_sig(ws.cell(c.row, c.column)) != style_sig(c): err(f'A-E {c.coordinate} beda dengan template')
    for col in 'ABCDE':
        if ws.column_dimensions[col].width != src.column_dimensions[col].width: err(f'lebar {col} beda template')
    hdr = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
    if hdr[-2:] != ['flag_mencurigakan', 'alasan_mencurigakan']: err('header flag salah')
    n = (len(hdr) - 7) // 3
    if hdr[5:5 + 3 * n] != sum([[f'nama_layanan_{k}', f'kelas_kamar_{k}', f'tarif_{k}'] for k in range(1, n + 1)], []): err('header triplet salah')
    if ws.freeze_panes != 'F2': err(f'freeze {ws.freeze_panes} != F2')
    from openpyxl.utils import get_column_letter as L
    for i in range(6, len(hdr) + 1):
        exp = [42, 24, 11][(i - 6) % 3] if i < 6 + 3 * n else (18 if i == 6 + 3 * n else 120)
        if ws.column_dimensions[L(i)].width != exp: err(f'lebar {L(i)} salah')
        h = ws.cell(1, i)
        if not h.font.b or not h.alignment.wrap_text: err(f'header {L(i)} tidak bold/wrap')
        if h.fill.fgColor.rgb[-6:] != ('FCE4D6' if i < 6 + 3 * n else 'FFC7CE'): err(f'fill header {L(i)} salah')
    FC, AC = 6 + 3 * n, 7 + 3 * n
    rows, titles = load_source(raw)
    byname = collections.defaultdict(list)
    for r in rows: byname[N_(r[2])].append(r)
    exact = {r[2] for r in rows}
    corpus = open(corpus_path).read()
    ex = list(openpyxl.load_workbook(excel, read_only=True).active.iter_rows(min_row=2, values_only=True))
    exidx = collections.defaultdict(list)
    for r in ex: exidx[(r[2], r[3])].append(N_(r[1]))
    groups = collections.defaultdict(list)
    allnames = set()
    for r in range(2, ws.max_row + 1):
        ura, kh = ws.cell(r, 2).value, ws.cell(r, 4).value
        groups[ura].append(r)
        fc, ac = ws.cell(r, FC), ws.cell(r, AC)
        if fc.value != 'YA' or not fc.font.b or fc.font.color is None or str(fc.font.color.rgb)[-6:] != '9C0006': err(f'flag baris {r} salah')
        if ac.alignment.vertical != 'top': err(f'alasan baris {r} tidak rata atas')
        for k in range(n):
            nm, lab, tr = (ws.cell(r, 6 + 3 * k + j) for j in range(3))
            if nm.value is None:
                if lab.value is not None or tr.value is not None: err(f'r{r} triplet {k+1} tidak lengkap')
                continue
            if nm.value == '-':
                if (lab.value, tr.value) != ('-', '-'): err(f'r{r} triplet "-" tidak lengkap')
                continue
            if not isinstance(tr.value, (int, float)) or tr.number_format != '#,##0': err(f'r{r} tarif {k+1} bukan angka #,##0')
            name, label, tarif = nm.value, lab.value, tr.value
            allnames.add(name)
            # label
            if label in KELAS.values():
                if label != KELAS[kh]: err(f'r{r} label {label} tidak sesuai baris {kh}')
            elif label not in EXTRA_LABELS and label != 'UMUM/NON-KELAS': err(f'r{r} label tidak diizinkan: {label}')
            if BAD_LABEL.search(str(label)) and label not in KELAS.values(): err(f'r{r} label kelas di atas I/ruang khusus: {label}')
            # nama + akhiran
            base, suffix = name, None
            m = re.match(r'^(.*) - ([A-K]\. [A-Z &]+)$', name)
            if m and m.group(1) in exact: base, suffix = m.group(1), m.group(2)
            if base not in exact: err(f'r{r} nama tidak ada persis di sumber: {name!r}'); continue
            occ = byname[N_(base)]
            same_b = N_(base) == N_(ura)
            if suffix:
                if same_b: err(f'r{r} akhiran pada nama = kolom B: {name}')
                if len(occ) < 2: err(f'r{r} akhiran pada nama unik: {name}')
                cand = [o for o in occ if o[1] == suffix and suffix in titles.get(o[0], []) + [o[1]]]
                if not cand: err(f'r{r} akhiran bukan judul di atas item: {name}'); continue
                occ = cand
            elif len(occ) > 1 and not same_b:
                if len({tuple(sorted(o[3].items())) for o in occ}) > 1 or len({o[1] for o in occ}) > 1:
                    err(f'r{r} nama kembar tanpa akhiran: {name}')
            # tarif
            col = label if label in KELAS.values() else ('ED' if label == 'ED' else 'RATES')
            if not any(o[3].get(col) == tarif for o in occ): err(f'r{r} tarif {tarif} tidak sama dengan sumber kolom {col} untuk {name}')
            if not any(N_(e) in N_(base) or N_(base) in N_(e) for e in exidx.get((label, tarif), []) if e): err(f'r{r} (label, tarif) tidak ada di Excel: {name} {label} {tarif}')
        # flag
        txt = ac.value or ''
        parts = txt.split(' | ')
        for i, p in enumerate(parts, 1):
            if not p.startswith(f'{i}. '): err(f'r{r} penomoran alasan salah di bagian {i}')
        for q in re.findall(r"'([^']+)'", txt):
            if q not in corpus: err(f'r{r} kutipan tidak ada di sumber: {q!r}')
    lower = collections.defaultdict(set)
    for nm in allnames: lower[nm.lower()].add(nm)
    for k, v in lower.items():
        if len(v) > 1: err(f'nama hanya beda huruf besar/kecil: {v}')
    for ura, rs in groups.items():
        if len(rs) != 3: err(f'{ura}: bukan 3 baris')
        if len({ws.cell(r, AC).value for r in rs}) != 1: err(f'{ura}: alasan beda antar kelas')
        seqs = {tuple(ws.cell(r, 6 + 3 * k).value for k in range(n)) for r in rs}
        if len(seqs) != 1: err(f'{ura}: urutan triplet beda antar kelas')
    return E

if __name__ == '__main__':
    e = check(*sys.argv[1:7])
    print('\n'.join(e) if e else 'RE-CHECK OK: 0 error')
