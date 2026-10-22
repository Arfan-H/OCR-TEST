import sys, shutil, copy, os
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import importlib

SRC, OUT, SPECMOD, SHEET = sys.argv[1:5]
m = importlib.import_module(SPECMOD); SPEC, NTRIP = m.SPEC, m.NTRIP
backup = os.path.join(os.path.dirname(OUT), 'backup_before_' + SHEET.replace(' ', '_') + '.xlsx')
shutil.copy(SRC, backup)
if os.path.abspath(SRC) != os.path.abspath(OUT): shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)
tpl = wb['Normalisasi']
pos = None
if SHEET in wb.sheetnames:
    pos = wb.sheetnames.index(SHEET); del wb[SHEET]
ws = wb.create_sheet(SHEET, pos) if pos is not None else wb.create_sheet(SHEET)

# A-E copy (value + style), widths, heights
for r in range(1, tpl.max_row + 1):
    for c in range(1, 6):
        s = tpl.cell(r, c); d = ws.cell(r, c, s.value)
        if s.has_style:
            d.font = copy.copy(s.font); d.fill = copy.copy(s.fill); d.border = copy.copy(s.border)
            d.alignment = copy.copy(s.alignment); d.number_format = s.number_format; d.protection = copy.copy(s.protection)
    if tpl.row_dimensions[r].height is not None:
        ws.row_dimensions[r].height = tpl.row_dimensions[r].height
for col in 'ABCDE':
    if col in tpl.column_dimensions:
        ws.column_dimensions[col].width = tpl.column_dimensions[col].width

hfill = PatternFill('solid', fgColor='FCE4D6'); ffill = PatternFill('solid', fgColor='FFC7CE')
col = 6
for n in range(1, NTRIP + 1):
    for name, w in ((f'nama_layanan_{n}', 42), (f'kelas_kamar_{n}', 24), (f'tarif_{n}', 11)):
        c = ws.cell(1, col, name); c.font = Font(bold=True); c.fill = hfill
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
        col += 1
FLAGC = col; ALC = col + 1
for cc, name, w in ((FLAGC, 'flag_mencurigakan', 18), (ALC, 'alasan_mencurigakan', 120)):
    c = ws.cell(1, cc, name); c.font = Font(bold=True); c.fill = ffill
    ws.column_dimensions[openpyxl.utils.get_column_letter(cc)].width = w

def flagged(text):
    keys = ['3. Ragu', '4. Janggal', 'Tidak ditemukan', 'selisih', 'non-kelas', 'mohon konfirmasi', 'tarif diisi -']
    return any(k in text for k in keys)

for r in range(2, tpl.max_row + 1):
    uraian = tpl.cell(r, 2).value; kelas = tpl.cell(r, 4).value
    sp = SPEC[uraian]
    col = 6
    for t in sp['triplets']:
        row = t['rows'][kelas]
        if sp.get('per_row_names'):
            name, label, tarif = row
        else:
            name = t['name']; label, tarif = row
        ws.cell(r, col, name)
        ws.cell(r, col + 1, label)
        c = ws.cell(r, col + 2, tarif)
        if isinstance(tarif, int): c.number_format = '#,##0'
        col += 3
    txt = sp['alasan']
    if flagged(txt):
        c = ws.cell(r, FLAGC, 'YA'); c.font = Font(bold=True, color='9C0006')
    a = ws.cell(r, ALC, txt); a.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'F2'
wb.save(OUT)
print('saved', OUT, 'backup', backup, 'sheets', wb.sheetnames)
