"""Uji mutasi generik: python3 mutate_gen.py OUT BAK PDF EX SHEET"""
import openpyxl, subprocess, sys, os, re, random
from openpyxl.styles import Font, PatternFill, Alignment
OUT, BAK, PDF, EX, S = sys.argv[1:6]
TMP = '/home/claude/work/mut/g.xlsx'; os.makedirs('/home/claude/work/mut', exist_ok=True)
CHK = '/home/claude/work/scripts/checker2.py'
wb0 = openpyxl.load_workbook(OUT); ws0 = wb0[S]; AC = ws0.max_column; FC = AC - 1
rows = list(range(2, ws0.max_row + 1)); random.seed(7)
M = []
def cellset(r, c, v):
    def f(wb): wb[S].cell(r, c).value = v
    return f
trip = [(r, c) for r in rows for c in range(6, FC, 3) if ws0.cell(r, c).value is not None]
for r, c in random.sample(trip, 12):
    nm = ws0.cell(r, c).value
    M.append((f'nama case r{r}', cellset(r, c, nm.swapcase())))
for r, c in random.sample(trip, 6):
    M.append((f'nama karangan r{r}', cellset(r, c, ws0.cell(r, c).value + ' X')))
for r, c in random.sample(trip, 12):
    t = ws0.cell(r, c + 2).value
    M.append((f'tarif r{r}', cellset(r, c + 2, (t + 1000) if isinstance(t, int) else 1000)))
labels = sorted({ws0.cell(r, c + 1).value for r, c in trip})
for r, c in random.sample(trip, 10):
    lb = ws0.cell(r, c + 1).value
    other = [l for l in labels if l != lb]
    M.append((f'label r{r}', cellset(r, c + 1, random.choice(other))))
for r in random.sample(rows, 5):
    t = ws0.cell(r, AC).value
    q = re.findall(r"'([^']+)'", t)[0]
    M.append((f'kutipan palsu r{r}', cellset(r, AC, t.replace(f"'{q}'", f"'{q} Palsu'", 1))))
    h = re.search(r'hlm\. (\d+)', t)
    if h: M.append((f'hlm r{r}', cellset(r, AC, t.replace(h.group(0), f'hlm. {int(h.group(1))+1}', 1))))
for r in random.sample(rows, 3):
    M.append((f'flag r{r}', cellset(r, FC, None if ws0.cell(r, FC).value == 'YA' else 'YA')))
def st(fn): return fn
M += [('A-E', cellset(5, 2, 'X')), ('qty', cellset(6, 3, 99)),
      ('normalisasi', lambda wb: setattr(wb['Normalisasi']['E7'], 'value', 1)),
      ('freeze', lambda wb: setattr(wb[S], 'freeze_panes', 'A1')),
      ('header', cellset(1, 7, 'kelas')),
      ('fill', lambda wb: setattr(wb[S].cell(1, 6), 'fill', PatternFill('solid', fgColor='FFFFFF'))),
      ('lebar', lambda wb: setattr(wb[S].column_dimensions['F'], 'width', 10)),
      ('wrap', lambda wb: setattr(wb[S].cell(3, AC), 'alignment', Alignment(wrap_text=False))),
      ('posisi', lambda wb: wb.move_sheet(S, offset=-1)),
      ('sheet ekstra', lambda wb: wb.create_sheet('Z'))]
for name in [s for s in wb0.sheetnames if s not in (S, 'Normalisasi')][:2]:
    M.append((f'sheet lain {name}', lambda wb, n=name: setattr(wb[n]['F2'], 'value', 'X')))
caught = 0; missed = []
for name, f in M:
    wb = openpyxl.load_workbook(OUT); f(wb); wb.save(TMP)
    p = subprocess.run([sys.executable, CHK, PDF, TMP, BAK, EX, S], capture_output=True, text=True)
    if p.returncode: caught += 1
    else: missed.append(name)
print(f'mutasi tertangkap: {caught}/{len(M)}'); print('lolos:', missed)
