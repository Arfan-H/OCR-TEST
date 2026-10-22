"""Checker generik (alat 2 = pdfplumber, dibatasi ke halaman yang dikutip) untuk satu sheet RS.
Usage: python3 checker2.py PDF OUT_XLSX BACKUP_XLSX EXCEL_RS SHEET"""
import sys, re, json, pickle, os, hashlib
import pdfplumber, pymupdf, openpyxl
from openpyxl.utils import get_column_letter as L

PDF, OUT, BAK, EXRS, SHEET = sys.argv[1:6]
ws_ = lambda s: re.sub(r'\s+', ' ', s or '').strip()
errors = []
err = errors.append

# ---------- KNOWN Excel differences (diverifikasi, tercatat di alasan) ----------
KNOWN = {
 'Mitra Keluarga Pondok Tjandra': {('UTAMA, Kelas I, ISO, PERINA', 'ISO/PERINA'), ('UTAMA, KELAS 1, ICU, ICCU, NICU, PICU, NINT, INT, ISO', 'ICU/ICCU/NICU/PICU/NINT/INT/ISO'), ('Kelas III, ODC', 'ODC')},
 'Mitra Keluarga Kelapa Gading': {('KELAS I / INT / IMC / IMA / NINT / PERINA / ISO', 'KELAS I/INT/IMC/PERINA/ISO')},
}.get(SHEET, set())
# alias nama Excel yang sudah diverifikasi (nama Excel tidak ada di PDF; tercatat di alasan)
NAME_ALIAS = {'Mitra Keluarga Jatiasih': {'KELAS 3': 'Kamar Perawatan – Kelas 3', 'KELAS 2': 'Kamar Perawatan – Kelas 2', 'KELAS 1': 'Kamar Perawatan – Kelas 1'}}.get(SHEET, {})

wb = openpyxl.load_workbook(OUT); bk = openpyxl.load_workbook(BAK)
if SHEET not in wb.sheetnames:
    print('sheet tidak ada'); sys.exit(1)
ws = wb[SHEET]
FC, AC = ws.max_column - 1, ws.max_column
alltext = ' '.join(str(ws.cell(r, AC).value or '') for r in range(2, ws.max_row + 1))
pages = sorted({int(p) for p in re.findall(r'hlm\. (\d+)', alltext)})

# ---------- PDF ----------
key = hashlib.md5((PDF + str(os.path.getsize(PDF)) + str(pages)).encode()).hexdigest()[:10]
cache = f'/tmp/claude-0/plumb_{key}.pkl'; os.makedirs('/tmp/claude-0', exist_ok=True)
if os.path.exists(cache):
    ROWS, CORP, NP = pickle.load(open(cache, 'rb'))
else:
    ROWS = []; CORP = set()
    doc = pymupdf.open(PDF); NP = len(doc)
    for p in doc:
        for l in p.get_text().split('\n'): CORP.add(ws_(l))
    with pdfplumber.open(PDF) as pdf:
        for pn in pages:
            if pn > len(pdf.pages): continue
            for t in pdf.pages[pn - 1].extract_tables():
                hdr = None; group = None
                for r in t:
                    c = [ws_(x) for x in r]
                    for x in c:
                        if x: CORP.add(x)
                    if any(x.startswith('KODE') for x in c) or ('EKSEKUTIF' in c and 'KELAS II' in c): hdr = c; hdr_isrow = not any(x.startswith('KODE') for x in c); continue
                    if hdr is None: continue
                    ki = [i for i, h in enumerate(hdr) if h.startswith('KODE')]
                    ki = ki[0] if ki else 0
                    nm_ = [i for i, h in enumerate(hdr) if h.startswith('NAMA')]
                    ni = nm_[0] if nm_ else ki + 1
                    pcols = [i for i, h in enumerate(hdr) if i > ni and not h.startswith('NAMA')]
                    vals = [c[i] if i < len(c) else '' for i in pcols]
                    if not any(re.search(r'\d', v) for v in vals):
                        ne = [x for x in c if x]
                        if ne: group = ne[-1]
                        continue
                    pr = {}
                    for i, v in zip(pcols, vals):
                        if re.fullmatch(r'(Rp)?[\d.,\s]+', v): v = v.replace(' ', '')
                        pr[hdr[i]] = v
                    ROWS.append(dict(page=pn, kode=c[ki], name=c[ni], group=group, prices=pr, header=[hdr[i] for i in pcols]))
    pickle.dump((ROWS, CORP, NP), open(cache, 'wb'))
CORP = list(CORP)
def in_pdf(q):
    rx = re.compile(r'(?<![A-Za-z0-9])' + re.escape(q) + r'(?![A-Za-z0-9])')
    return any(q in s and rx.search(s) for s in CORP)
def toint(v):
    v = re.sub(r'^Rp\s*', '', (v or '').strip())
    if re.fullmatch(r'\d{1,3}([.,]\d{3})*|\d+', v): return int(re.sub(r'[.,]', '', v))
    return None
def find(kode, name, grp=None):
    return [r for r in ROWS if r['name'] == name and (r['kode'] == kode or (r['kode'] == '' and grp and r['group'] == grp))]

# ---------- workbook structure ----------
def sig(c):
    return (c.value, bool(c.font.b), bool(c.font.i), c.font.color.rgb if c.font.color else None, c.font.name, c.font.sz,
            c.fill.fgColor.rgb, c.fill.fill_type, c.number_format, c.alignment.horizontal, c.alignment.vertical,
            bool(c.alignment.wrap_text), c.border.left.style, c.border.right.style, c.border.top.style, c.border.bottom.style)
others_b = [s for s in bk.sheetnames if s != SHEET]; others_o = [s for s in wb.sheetnames if s != SHEET]
if others_b != others_o: err('daftar sheet lain berubah')
exp_pos = bk.sheetnames.index(SHEET) if SHEET in bk.sheetnames else len(wb.sheetnames) - 1
if wb.sheetnames.index(SHEET) != exp_pos: err('posisi sheet salah')
for name in others_b:
    if name not in wb.sheetnames: continue
    a, b = wb[name], bk[name]
    if (a.max_row, a.max_column) != (b.max_row, b.max_column): err(f'sheet {name} dimensi berubah')
    for row in b.iter_rows():
        for c in row:
            if sig(c) != sig(a[c.coordinate]): err(f'sheet {name} sel {c.coordinate} berubah'); break
    for k, d in b.column_dimensions.items():
        if a.column_dimensions[k].width != d.width: err(f'sheet {name} lebar {k} berubah')
    if a.freeze_panes != b.freeze_panes: err(f'sheet {name} freeze berubah')
tpl = bk['Normalisasi']
for r in range(1, tpl.max_row + 1):
    for c in range(1, 6):
        if sig(tpl.cell(r, c)) != sig(ws.cell(r, c)): err(f'A-E beda {ws.cell(r, c).coordinate}')
    if tpl.row_dimensions[r].height != ws.row_dimensions[r].height: err(f'tinggi baris {r}')
for k in 'ABCDE':
    if tpl.column_dimensions[k].width != ws.column_dimensions[k].width: err(f'lebar {k}')
if ws.max_row != tpl.max_row: err('jumlah baris beda')
if ws.freeze_panes != 'F2': err('freeze bukan F2')
hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
ntr = (len(hdr) - 7) // 3
if hdr != hdr[:5] + [f'{p}_{n}' for n in range(1, ntr + 1) for p in ('nama_layanan', 'kelas_kamar', 'tarif')] + ['flag_mencurigakan', 'alasan_mencurigakan']:
    err(f'header salah')
for i in range(6, 6 + 3 * ntr):
    c = ws.cell(1, i)
    if not c.font.b or c.fill.fgColor.rgb[-6:] != 'FCE4D6': err(f'style header {c.coordinate}')
    if ws.column_dimensions[L(i)].width != [42, 24, 11][(i - 6) % 3]: err(f'lebar {L(i)}')
for i, w in ((FC, 18), (AC, 120)):
    c = ws.cell(1, i)
    if not c.font.b or c.fill.fgColor.rgb[-6:] != 'FFC7CE': err(f'style header {c.coordinate}')
    if ws.column_dimensions[L(i)].width != w: err(f'lebar {L(i)}')

# ---------- Excel RS ----------
EXN = {}; EXK = set(); EXL = set(); EXNM = set()
for r in openpyxl.load_workbook(EXRS, read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    if r[1] is None: continue
    EXN.setdefault(ws_(str(r[1])).lower(), set()).add((r[2], r[3])); EXK.add(r[0]); EXL.add(r[2]); EXNM.add(ws_(str(r[1])))

TITLES = {1: ('Nama sama', 'Nama beda', 'Tidak ditemukan'), 2: ('Pemetaan kelas & label',), 3: ('Ragu',), 4: ('Janggal',),
          5: ('Excel vs sumber',), 6: ('Catatan sumber',), 7: ('Tidak diambil',), 8: ('Penamaan',)}
CLS = {'III': re.compile(r'Kelas III|KELAS III|(^|[/ ])III([/ ]|$)|KELAS 3|Kelas 3'), 'II': re.compile(r'Kelas II\b(?!I)|KELAS II\b(?!I)|^II$|KELAS 2|Kelas 2'),
       'I & R. KHUSUS': re.compile(r'Kelas 1|Kelas I\b(?!I)|KELAS 1|KELAS I\b(?!I)')}
TOK = {'III': {'III'}, 'II': {'II'}, 'I & R. KHUSUS': {'I', '1'}}
def extoks(lbl): return set(re.split(r'[/ ]+', (lbl or '').replace('KELAS', '').strip()))
uraian_rows = {}
for r in range(2, ws.max_row + 1): uraian_rows.setdefault(ws.cell(r, 2).value, []).append(r)
for ur, rows in uraian_rows.items():
    texts = [ws.cell(r, AC).value for r in rows]
    if len(set(texts)) != 1: err(f'{ur}: alasan beda antar baris')
    text = texts[0] or ''
    cites = re.findall(r"'([^']+)' \(kode ([A-Z0-9]{10})(?:, (?:kelompok|tabel) '([^']+)')?", text)
    for name, kode, hl in re.findall(r"'([^']+)' \(kode ([A-Z0-9]{10})[^)]*?hlm\. (\d+)", text):
        g = re.search(re.escape(f"'{name}' (kode {kode}, kelompok '") + r"([^']+)'", text)
        rr = find(kode, name, g.group(1) if g else None)
        if not rr: err(f'{ur}: kode {kode} {name!r} tidak ada di halaman PDF yang dikutip')
        elif str(rr[0]['page']) != hl: err(f'{ur}: hlm salah untuk {kode}')
    pts = text.split(' | '); nums = []
    for p in pts:
        m = re.match(r'(\d)\. ([^:]+):', p)
        if not m: err(f'{ur}: poin tanpa format'); continue
        n = int(m.group(1)); nums.append(n)
        if not any(m.group(2).startswith(t) for t in TITLES.get(n, ('#',))): err(f'{ur}: judul poin {n} salah')
    if nums != sorted(set(nums)): err(f'{ur}: penomoran tidak urut')
    for q in re.findall(r"'([^']+)'", text):
        if not in_pdf(q): err(f'{ur}: kutipan tunggal tidak ada di PDF: {q!r}')
    for q in re.findall(r'"([^"]+)"', text):
        if not (q in EXK or q in EXL or q in EXNM or q.startswith(' - ')): err(f'{ur}: kutipan ganda tidak dikenal: {q!r}')
    need = any(k in text for k in ('3. Ragu', '4. Janggal', 'Tidak ditemukan', 'selisih', 'non-kelas', 'mohon konfirmasi', 'tarif diisi -'))
    for r in rows:
        f = ws.cell(r, FC)
        if need != (f.value == 'YA'): err(f'{ur} r{r}: flag salah')
        if f.value == 'YA' and (not f.font.b or (f.font.color.rgb or '')[-6:] != '9C0006'): err(f'{ur} r{r}: style flag')
        a = ws.cell(r, AC)
        if not a.alignment.wrap_text or a.alignment.vertical != 'top': err(f'{ur} r{r}: style alasan')
    per_row = 'tiap baris memakai item kelasnya sendiri' in text
    seqs = []
    for r in rows:
        kelas = ws.cell(r, 4).value; seq = []
        for t in range(ntr):
            nm, lb, tf = (ws.cell(r, 6 + 3 * t + k).value for k in range(3))
            if nm is None and lb is None and tf is None: seq.append(None); continue
            if None in (nm, lb, tf): err(f'{ur} r{r} t{t+1}: triplet tidak lengkap'); seq.append('?'); continue
            if isinstance(tf, int) and ws.cell(r, 8 + 3 * t).number_format != '#,##0': err(f'{ur} r{r}: format tarif')
            row = None; base = nm
            ck = [(k, g) for n_, k, g in cites if n_ == nm]
            if ck:
                rr = [x for k, g in ck for x in find(k, nm, g)]
                if len({k for k, g in ck}) > 1: err(f'{ur}: nama {nm!r} dikutip >1 kode tanpa akhiran')
                elif len(rr) >= 1: row = rr[0]
            elif ' - ' in nm:
                b, s = nm.rsplit(' - ', 1)
                ck2 = [k for n_, k, g in cites if n_ == b and g == s]
                if len(ck2) == 1:
                    rr = find(ck2[0], b)
                    if rr: row = rr[0]; base = b
                    if sum(1 for n_, k, g in cites if n_ == b) < 2: err(f'{ur}: akhiran pada nama yang hanya sekali dimasukkan {nm!r}')
            if row is None: err(f'{ur} r{r}: nama {nm!r} tidak ditemukan persis di PDF'); seq.append(nm); continue
            filled = [h for h in row['header'] if toint(row['prices'][h]) is not None]
            if lb in row['header']:
                v = row['prices'][lb]
                if tf == '-':
                    if toint(v) is not None: err(f'{ur} r{r}: tarif "-" padahal sel PDF berisi')
                elif toint(v) != tf: err(f'{ur} r{r}: tarif {tf} != PDF {v} ({nm}/{lb})')
                if not CLS[kelas].search(lb): err(f'{ur} r{r}: label {lb!r} tidak sesuai kelas {kelas}')
                exlabel = None
            else:  # label mengikuti Excel: hanya untuk item satu tarif
                if len(filled) != 1: err(f'{ur} r{r}: label {lb!r} bukan header PDF padahal item punya >1 kolom tarif')
                elif toint(row['prices'][filled[0]]) != tf: err(f'{ur} r{r}: tarif {tf} != PDF {row["prices"][filled[0]]}')
                if lb == 'UMUM/NON-KELAS':
                    if 'non-kelas' not in text: err(f'{ur}: non-kelas tanpa catatan')
                elif not (per_row and CLS[kelas].search(lb)): err(f'{ur} r{r}: label Excel {lb!r} tidak sah')
                exlabel = lb
            # Excel cross-check
            exs = EXN.get(NAME_ALIAS.get(base, base).lower(), set())
            if tf == '-': pass
            elif exlabel:
                if (exlabel, tf) not in exs: err(f'{ur} r{r}: Excel tidak punya ({exlabel},{tf}) untuk {base!r}')
            else:
                ok = any(v == tf and (TOK[kelas] & extoks(l) or (lb, l) in KNOWN) for l, v in exs)
                if not ok: err(f'{ur} r{r}: Excel tidak cocok untuk {base!r} {lb!r} {tf}')
            seq.append(re.sub(r'(?i)\s*-?\s*kelas\s*(III|II|I|1|2|3)\b.*$', '', nm) if per_row else nm)
        seqs.append(seq)
    if len(set(map(tuple, seqs))) != 1: err(f'{ur}: urutan/daftar triplet beda antar baris')
    lows = [n.lower() for n in seqs[0] if n]
    if len(lows) != len(set(lows)): err(f'{ur}: nama kembar/beda kapital')
print(f'ERRORS: {len(errors)}' + ''.join('\n - ' + e for e in errors[:40]))
sys.exit(1 if errors else 0)
