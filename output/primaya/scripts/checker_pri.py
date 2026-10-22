"""Checker PRIMAYA (turunan checker_mika.py; alat 2 = kata PyMuPDF + pembaca posisi linegrid, terpisah dari ekstraksi pdfplumber).
Usage: python3 checker_pri.py PDF OUT_XLSX BASE_XLSX EXCEL_RS SHEET
BASE_XLSX = workbook sebelum sheet-sheet baru ditambahkan (sheet yang ada di BASE, selain SHEET, wajib tidak berubah).
Yang dicek: nama persis di PDF (kode + hlm. yang dikutip), tarif per label kolom, label kelas, kutipan '...' ada di PDF,
kutipan "..." ada di Excel, (label, tarif) ada di Excel, format poin alasan, flag, style, struktur sheet."""
import sys, re, os, pickle, hashlib
import pymupdf, openpyxl
from openpyxl.utils import get_column_letter as L

PDF, OUT, BASE, EXRS, SHEET = sys.argv[1:6]
ws_ = lambda s: re.sub(r'\s+', ' ', str(s or '')).strip()
errors = []; err = errors.append
KODE = re.compile(r'^[A-Z]{4}\d{6}$')


def toint(v):
    v = re.sub(r'^Rp\.?\s*', '', ws_(v))
    v = re.sub(r'(?<=\d{3})[.,]00$', '', v) if re.fullmatch(r'\d{1,3}([.,]\d{3})+[.,]00', v) else v   # '693.000,00'
    if re.fullmatch(r'\d{1,3}([.,]\d{3})+|\d+', v): return int(re.sub(r'[.,]', '', v))
    return None


def ctok(lbl):
    """token kelas reguler (I/II/III) dalam label; angka 1/2/3 setara romawi"""
    out = set()
    for t in re.split(r'[,/&\s]+', ws_(lbl).upper().replace('*', '')):
        t = {'1': 'I', '2': 'II', '3': 'III'}.get(t.strip('. '), t.strip('. '))
        if t in ('I', 'II', 'III'): out.add(t)
    return out


WANT = {'III': 'III', 'II': 'II', 'I & R. KHUSUS': 'I'}

wb = openpyxl.load_workbook(OUT); bk = openpyxl.load_workbook(BASE)
if SHEET not in wb.sheetnames: print('sheet tidak ada'); sys.exit(1)
ws = wb[SHEET]
FC, AC = ws.max_column - 1, ws.max_column
alltext = ' '.join(str(ws.cell(r, AC).value or '') for r in range(2, ws.max_row + 1))
pages = sorted({int(p) for p in re.findall(r'hlm\. (\d+)', alltext)})

class _Row:
    def __init__(self, cells):
        self.cells = cells
        bb = [c for c in cells if c]
        self.bbox = (min(c[0] for c in bb), min(c[1] for c in bb), max(c[2] for c in bb), max(c[3] for c in bb))


class _Tab:
    def __init__(self, rows): self.rows = rows


def grid_tables(page):
    """Tabel dari garis yang digambar (dipakai di halaman ber-watermark, tempat find_tables PyMuPDF gagal memisah baris).
    Kolom = garis vertikal yang menutupi baris itu; kolom tanpa garis pemisah = sel gabung (None), seperti find_tables."""
    V = []; H = []
    for dr in page.get_drawings():
        for it in dr['items']:
            if it[0] == 're':
                r = it[1]
                if r.width < 1.5 and r.height > 4: V.append(((r.x0 + r.x1) / 2, r.y0, r.y1))
                elif r.height < 1.5 and r.width > 4: H.append(((r.y0 + r.y1) / 2, r.x0, r.x1))
            elif it[0] == 'l':
                a, b = it[1], it[2]
                if abs(a.x - b.x) < 0.5 and abs(a.y - b.y) > 4: V.append((a.x, min(a.y, b.y), max(a.y, b.y)))
                elif abs(a.y - b.y) < 0.5 and abs(a.x - b.x) > 4: H.append((a.y, min(a.x, b.x), max(a.x, b.x)))
    # kelompokkan garis vertikal menjadi tabel: rentang y yang bersambung/bertumpuk
    V.sort(key=lambda v: v[1]); groups = []
    for v in V:
        for g in groups:
            if v[1] <= g['y1'] + 1.5 and v[2] >= g['y0'] - 1.5: g['v'].append(v); g['y0'] = min(g['y0'], v[1]); g['y1'] = max(g['y1'], v[2]); break
        else: groups.append(dict(y0=v[1], y1=v[2], v=[v]))
    tabs = []
    for g in groups:
        xs = []
        for x in sorted(v[0] for v in g['v']):
            if not xs or x - xs[-1] > 2: xs.append(x)
        if len(xs) < 3: continue
        ys = []
        for y in sorted(h[0] for h in H if g['y0'] - 1.5 <= h[0] <= g['y1'] + 1.5 and h[1] <= xs[0] + 3 and h[2] >= xs[1] - 3):   # garis baris minimal melintasi kolom pertama (kolom lain bisa sel gabung)
            if not ys or y - ys[-1] > 2: ys.append(y)
        rows = []
        for y0, y1 in zip(ys, ys[1:]):
            ym = (y0 + y1) / 2
            have = [any(abs(v[0] - x) <= 2 and v[1] - 1 <= ym <= v[2] + 1 for v in g['v']) for x in xs]
            cells = []; start = 0
            for i in range(1, len(xs)):
                if have[i] or i == len(xs) - 1:
                    cells.append((xs[start], y0, xs[i], y1)); cells += [None] * (i - start - 1); start = i
            rows.append(_Row(cells))
        if rows: tabs.append(_Tab(rows))
    return tabs


# ---------- PDF (PyMuPDF, pembaca posisi kata linegrid) ----------
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from linegrid import read_page
# teks PDF apa adanya (tanpa dipotong clip path sel): nama yang meluber dan terpotong tampilan sel tetap terbaca utuh
FL = pymupdf.TEXTFLAGS_WORDS & ~pymupdf.TEXT_MEDIABOX_CLIP
_lg = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'linegrid.py')).read()   # cache ikut versi linegrid
key = hashlib.md5(('pri6' + _lg + os.path.abspath(PDF) + str(os.path.getsize(PDF)) + str(pages)).encode()).hexdigest()[:12]
cache = os.path.join(os.environ.get('CHK_CACHE', '/tmp'), f'mupdf_{key}.pkl')
if os.path.exists(cache):
    ROWS, CORP, PTXT = pickle.load(open(cache, 'rb'))
else:
    ROWS = []; CORP = set()
    doc = pymupdf.open(PDF)
    PTXT = [re.sub(r'\s+', '', p.get_text(flags=FL)) for p in doc]   # teks per halaman tanpa spasi (cek judul/kelompok per halaman)
    for p in doc:
        for l in p.get_text(flags=FL).split('\n'): CORP.add(ws_(l))
        CORP.add(ws_(p.get_text(flags=FL)))           # nama terlipat ke beberapa baris
    prev = None
    for pn in range(1, (max(pages) if pages else 0) + 1):   # berurutan: halaman lanjutan memakai header halaman sebelumnya
        if pn > len(doc): break
        words = [w[:5] for w in doc[pn - 1].get_text('words', flags=FL)]
        items, prev = read_page(words, prev)
        if pn not in pages: continue
        for it in items:
            ROWS.append(dict(page=pn, kode=it['kode'], merged=False, cont=False, name=ws_(it['name']), prices=it['prices'], near=[],
                             cells=[it['name']] + list(it['prices'].values()), main={k: k for k in it['prices']}))
    pickle.dump((ROWS, CORP, PTXT), open(cache, 'wb'))
CORP = list(CORP)


CORPW = {re.sub(r'\s+', '', s) for s in CORP if len(s) > 200}   # teks halaman tanpa spasi (label header terlipat beberapa baris)


def in_pdf(q):
    rx = re.compile(r'(?<![A-Za-z0-9])' + re.escape(q) + r'(?![A-Za-z0-9])')
    if any(q in s and rx.search(s) for s in CORP): return True
    qw = re.sub(r'\s+', '', q)
    return len(qw) >= 6 and any(qw in s for s in CORPW)


NZ = lambda s: re.sub(r'[^A-Z0-9]', '', (s or '').upper())


def col_of(row, lb):
    """kolom PDF untuk label sheet: sama persis, atau potongan berurutan dari gabungan header kolom
    yang memuat teks baris header utama kolom itu (header multi-baris / label rentang di atasnya)"""
    if lb in row['prices']: return lb
    W = lambda s: re.sub(r'\s+', '', s or '')   # teks label harus persis (abaikan spasi/baris), tanda baca & kapital dihitung
    m = [k for k in row['prices'] if NZ(lb) and NZ(lb) in NZ(k) and NZ(row['main'].get(k)) and NZ(row['main'].get(k)) in NZ(lb) and W(lb) in W(k)]
    return m[0] if len(m) == 1 else None


BUL = lambda s: re.sub(r'^- +', '', s or '')     # tanda daftar '- ' di depan nama (wajib dicatat di poin 8)


def find(kode, name, page, rank=False):
    m = [r for r in ROWS if r['kode'] == kode and BUL(r['name']) == name and r['page'] == page and not r.get('cont')]
    if rank:    # sub-baris (kode & nama rowspan) hanya dipakai bila dikutip dengan isi kolom pembeda
        m += [r for r in ROWS if r.get('cont') and r['kode'] == kode and BUL(r['name']) == name and r['page'] == page and kode in r['near']]
    if not m:   # sel kode gabung tanpa kode di baris ini: kode harus tertulis di dekat baris (<40pt)
        m = [r for r in ROWS if r['kode'] == '' and r['merged'] and BUL(r['name']) == name and r['page'] == page and kode in r['near']]
    return m


# ---------- struktur workbook ----------
def sig(c):
    return (c.value, bool(c.font.b), bool(c.font.i), c.font.color.rgb if c.font.color else None, c.font.name, c.font.sz,
            c.fill.fgColor.rgb, c.fill.fill_type, c.number_format, c.alignment.horizontal, c.alignment.vertical,
            bool(c.alignment.wrap_text), c.border.left.style, c.border.right.style, c.border.top.style, c.border.bottom.style)


if wb.sheetnames[:len(bk.sheetnames)] != bk.sheetnames and SHEET not in bk.sheetnames: err('urutan sheet BASE berubah')
for name in bk.sheetnames:
    if name == SHEET: continue
    if name not in wb.sheetnames: err(f'sheet {name} hilang'); continue
    a, b = wb[name], bk[name]
    if (a.max_row, a.max_column) != (b.max_row, b.max_column): err(f'sheet {name} dimensi berubah')
    for row in b.iter_rows():
        for c in row:
            if sig(c) != sig(a[c.coordinate]): err(f'sheet {name} sel {c.coordinate} berubah'); break
    for k, d in b.column_dimensions.items():
        if a.column_dimensions[k].width != d.width: err(f'sheet {name} lebar {k} berubah')
    if a.freeze_panes != b.freeze_panes: err(f'sheet {name} freeze berubah')
tpl = bk['Normalisasi']
for r in range(1, tpl.max_row + 1):
    for c in range(1, 6):
        if sig(tpl.cell(r, c)) != sig(ws.cell(r, c)): err(f'A-E beda {ws.cell(r, c).coordinate}')
    if tpl.row_dimensions[r].height != ws.row_dimensions[r].height: err(f'tinggi baris {r}')
for k in 'ABCDE':
    if tpl.column_dimensions[k].width != ws.column_dimensions[k].width: err(f'lebar {k}')
if ws.max_row != tpl.max_row: err('jumlah baris beda')
if ws.freeze_panes != 'F2': err('freeze bukan F2')
hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
ntr = (len(hdr) - 7) // 3
if hdr != [tpl.cell(1, c).value for c in range(1, 6)] + [f'{p}_{n}' for n in range(1, ntr + 1) for p in ('nama_layanan', 'kelas_kamar', 'tarif')] + ['flag_mencurigakan', 'alasan_mencurigakan']:
    err('header salah')
for i in range(6, 6 + 3 * ntr):
    c = ws.cell(1, i)
    if not c.font.b or (c.fill.fgColor.rgb or '')[-6:] != 'FCE4D6': err(f'style header {c.coordinate}')
    if ws.column_dimensions[L(i)].width != [42, 24, 11][(i - 6) % 3]: err(f'lebar {L(i)}')
for i, w in ((FC, 18), (AC, 120)):
    c = ws.cell(1, i)
    if not c.font.b or (c.fill.fgColor.rgb or '')[-6:] != 'FFC7CE': err(f'style header {c.coordinate}')
    if ws.column_dimensions[L(i)].width != w: err(f'lebar {L(i)}')

# ---------- Excel RS ----------
exws = openpyxl.load_workbook(EXRS, read_only=True).worksheets[0]
it = exws.iter_rows(values_only=True); eh = next(it)
ik, inm, ikl, itf = (eh.index(k) for k in ('kategori_harga', 'nama_layanan', 'kelas_kamar', 'tarif'))
EXN = {}; EXSTR = set()
for r in it:
    if r[inm] is None: continue
    n = ws_(r[inm])
    EXN.setdefault(n.lower(), set()).add((ws_(r[ikl]), r[itf]))
    EXSTR.update({n, ws_(r[ikl]), ws_(r[ik])})

TITLES = {1: ('Nama sama', 'Nama beda', 'Tidak ditemukan'), 2: ('Pemetaan kelas & label',), 3: ('Ragu',), 4: ('Janggal',),
          5: ('Excel vs sumber',), 6: ('Catatan sumber',), 7: ('Tidak diambil',), 8: ('Penamaan',)}
FLAGK = ('3. Ragu', '4. Janggal', 'Tidak ditemukan', 'selisih', 'non-kelas', 'mohon konfirmasi', 'tarif diisi -')
RANKCOL = ('Rank', 'RANK', 'KETERANGAN', 'JENIS KAMAR', 'KELAS')
CITE = re.compile(r"'([^']+)' \((?:kode ([A-Z0-9]+)|tanpa kode)( sel gabung)?((?:, (?:kelompok|tabel|Rank|RANK|KETERANGAN|JENIS KAMAR|KELAS) '[^']+')*), hlm\. (\d+)")
uraian_rows = {}
for r in range(2, ws.max_row + 1): uraian_rows.setdefault(ws.cell(r, 2).value, []).append(r)
if set(uraian_rows) != {tpl.cell(r, 2).value for r in range(2, tpl.max_row + 1)}: err('daftar uraian beda')
for ur, rows in uraian_rows.items():
    texts = [ws.cell(r, AC).value for r in rows]
    if len(set(texts)) != 1: err(f'{ur}: alasan beda antar baris')
    text = texts[0] or ''
    cites = []; used = {}
    for name, kode, mg, extra, hl in CITE.findall(text):
        g = re.search(r"kelompok '([^']+)'", extra); t = re.search(r"tabel '([^']+)'", extra)
        rk = re.search(r"(Rank|RANK|KETERANGAN|JENIS KAMAR|KELAS) '([^']+)'", extra)
        c = dict(name=name, kode=kode, page=int(hl), group=g.group(1) if g else None, title=t.group(1) if t else None,
                 rank=rk.group(2) if rk else None, row=None)
        fr = find(kode, name, int(hl), bool(rk))
        if rk: fr = [r for r in fr if c['rank'] in r['cells']]   # isi kolom pembeda (Rank/KETERANGAN) ada di baris itu
        k = (kode, name, int(hl), c['rank']); i = used.get(k, 0); used[k] = i + 1
        if i >= len(fr): err(f'{ur}: {kode} {name!r} tidak ada di hlm. {hl} (PyMuPDF)')
        else:
            c['row'] = fr[i]; c['nrows'] = len(fr)
            if bool(mg) != fr[i]['merged']: err(f'{ur}: status sel gabung {kode} {name!r} salah')
            if fr[i]['name'] != name and "tanda daftar '- '" not in text: err(f'{ur}: {name!r} di PDF berawalan tanda daftar tetapi tidak dicatat')
        cites.append(c)
    # judul tabel/kelompok yang dikutip harus tercetak di halaman item, atau dicatat 'tercetak di hlm. N' dan memang ada di hlm. N
    carried = {(k, t, int(b)): int(a) for k, t, a, b in re.findall(r"(judul tabel|kelompok) '([^']+)' tercetak di hlm\. (\d+) \(tabel berlanjut ke hlm\. (\d+)\)", text)}
    onpg = lambda t, pg: 0 < pg <= len(PTXT) and re.sub(r'\s+', '', t) in PTXT[pg - 1]
    for (k, t, b), a in carried.items():
        if not (a < b and onpg(t, a)): err(f'{ur}: {k} {t!r} tidak tercetak di hlm. {a} (dicatat untuk hlm. {b})')
    for c in cites:
        for k, key in (('judul tabel', 'title'), ('kelompok', 'group')):
            t = c[key]
            if t and not onpg(t, c['page']) and (k, t, c['page']) not in carried:
                err(f"{ur}: {k} {t!r} tidak tercetak di hlm. {c['page']} dan tanpa catatan halaman judul")
    pts = text.split(' | '); nums = []
    for p in pts:
        m = re.match(r'(\d)\. ([^:]+):', p)
        if not m: err(f'{ur}: poin tanpa format'); continue
        n = int(m.group(1)); nums.append(n)
        if not any(m.group(2).startswith(t) for t in TITLES.get(n, ('#',))): err(f'{ur}: judul poin {n} salah')
    if nums != sorted(set(nums)) or nums[:2] != [1, 2] or 5 not in nums or 8 not in nums: err(f'{ur}: penomoran poin salah')
    for q in re.findall(r"'([^']+)'", text):
        if not in_pdf(q): err(f'{ur}: kutipan tunggal tidak ada di PDF: {q!r}')
    for q in re.findall(r'"([^"]+)"', text):
        if not (q in EXSTR or q.startswith(' - ') or f'"{q}"' in EXSTR or in_pdf(f'"{q}"')): err(f'{ur}: kutipan ganda tidak ada di Excel: {q!r}')
    alias = {}
    for e, p, rv in re.findall(r"Excel menamai item ini \"([^\"]+)\" \(PDF '([^']+)'(?: (?:Rank|RANK|KETERANGAN|JENIS KAMAR|KELAS) '([^']+)')?\)", text):
        alias[(p, rv or None)] = e
    need = any(k in text for k in FLAGK)
    for r in rows:
        f = ws.cell(r, FC)
        if need != (f.value == 'YA'): err(f'{ur} r{r}: flag salah')
        if f.value == 'YA' and (not f.font.b or (f.font.color.rgb or '')[-6:] != '9C0006'): err(f'{ur} r{r}: style flag')
        a = ws.cell(r, AC)
        if not a.alignment.wrap_text or a.alignment.vertical != 'top': err(f'{ur} r{r}: style alasan')
    per_row = 'tiap baris memakai item kelasnya sendiri' in text
    nf = text.startswith('1. Tidak ditemukan:')
    if nf:
        # istilah yang dinyatakan tidak ada harus benar-benar tidak ada di teks PDF (semua halaman) dan di nama layanan Excel
        terms = re.findall(r'tidak ada teks «([^»]+)» di PDF', text); xterms = re.findall(r'Excel tidak memuat nama berisi «([^»]+)»', text)
        if not terms or not xterms: err(f'{ur}: Tidak ditemukan tanpa istilah pencarian PDF/Excel')
        LOW = [c.lower() for c in CORP]
        for q in terms:
            if any(q.lower() in c for c in LOW) or any(re.sub(r'\s+', '', q.lower()) in re.sub(r'\s+', '', c) for c in LOW): err(f'{ur}: {q!r} ternyata ada di PDF')
        for q in xterms:
            if any(q.lower() in n for n in EXN): err(f'{ur}: {q!r} ternyata ada di Excel')
    seqs = []
    for r in rows:
        kelas = ws.cell(r, 4).value; seq = []
        for t in range(ntr):
            nm, lb, tf = (ws.cell(r, 6 + 3 * t + k).value for k in range(3))
            if nm is None and lb is None and tf is None: seq.append(None); continue
            if (nm, lb, tf) == ('-', '-', '-'):
                if not nf or t: err(f'{ur} r{r}: triplet - tanpa Tidak ditemukan')
                seq.append('-'); continue
            if nf: err(f'{ur} r{r}: Tidak ditemukan tetapi triplet berisi')
            if None in (nm, lb, tf): err(f'{ur} r{r} t{t + 1}: triplet tidak lengkap'); seq.append('?'); continue
            if isinstance(tf, int) and ws.cell(r, 8 + 3 * t).number_format != '#,##0': err(f'{ur} r{r}: format tarif')
            # resolusi nama -> kutipan (baris PDF sudah dipetakan per kutipan)
            cand = [c for c in cites if c['name'] == nm]; base = nm; row = None
            if cand:
                if len(cand) > 1: err(f'{ur}: nama {nm!r} dikutip >1x tanpa akhiran')
                elif cand[0]['row'] and cand[0]['nrows'] > 1 and not cand[0]['rank']: err(f'{ur}: {nm!r} >1 baris di hlm. {cand[0]["page"]} tetapi tanpa akhiran')
                row = cand[0]['row']
            elif ' - ' in nm:
                b, s = nm.rsplit(' - ', 1); base = b
                same = [c for c in cites if c['name'] == b]
                if len(same) < 2: err(f'{ur}: akhiran pada nama yang hanya sekali dimasukkan {nm!r}')
                m = re.fullmatch(r'baris (\d+)', s)
                if m:
                    i = int(m.group(1)) - 1
                    if i < len(same) and len({(c['kode'], c['page']) for c in same}) == 1: row = same[i]['row']
                else:
                    c2 = [c for c in same if s in (c['group'], c['title'], c['rank'], c['kode'])]
                    if len(c2) == 1: row = c2[0]['row']
            if row is None: err(f'{ur} r{r}: nama {nm!r} tidak ditemukan persis di PDF'); seq.append(nm); continue
            filled = [h for h in row['prices'] if toint(row['prices'][h]) is not None]
            ck = col_of(row, lb); nonhdr = False
            if ck is not None:
                v = row['prices'][ck]
                if tf == '-':
                    if toint(v) is not None: err(f'{ur} r{r}: tarif "-" padahal sel PDF berisi')
                    if 'tarif diisi -' not in text: err(f'{ur}: tarif - tanpa catatan')
                elif toint(v) != tf: err(f'{ur} r{r}: tarif {tf} != PDF {v} ({nm}/{lb})')
                exlabel = None
                if WANT[kelas] not in ctok(lb):
                    if filled == [ck] and 'non-kelas mengikuti Excel' in text: exlabel = lb
                    else: err(f'{ur} r{r}: label {lb!r} tidak sesuai kelas {kelas}')
            else:
                if len(filled) != 1: err(f'{ur} r{r}: label {lb!r} bukan header PDF padahal item punya {len(filled)} kolom tarif')
                elif toint(row['prices'][filled[0]]) != tf: err(f'{ur} r{r}: tarif {tf} != PDF {row["prices"][filled[0]]}')
                if lb == 'UMUM/NON-KELAS':
                    if 'non-kelas' not in text: err(f'{ur}: non-kelas tanpa catatan')
                elif not (per_row and WANT[kelas] in ctok(lb)) and 'non-kelas mengikuti Excel' not in text \
                        and not (per_row and f'label mengikuti Excel ("{lb}"' in text and f'Excel memberi label "{lb}"' in text and '4. Janggal' in text):
                    err(f'{ur} r{r}: label Excel {lb!r} tidak sah')
                exlabel = lb; nonhdr = True
            rv = next((c['rank'] for c in cites if c['row'] is row), None)
            exs = EXN.get(alias.get((base, rv), base).lower(), set())
            if not exs:
                if f"'{base}' tidak ada di Excel" not in text: err(f'{ur}: {base!r} tidak ada di Excel tanpa catatan')
                if nonhdr and lb != 'UMUM/NON-KELAS': err(f'{ur} r{r}: label {lb!r} bukan header PDF dan item tidak ada di Excel')
            elif tf == '-': pass
            elif exlabel:
                if (exlabel, tf) not in exs: err(f'{ur} r{r}: Excel tidak punya ({exlabel},{tf}) untuk {base!r}')
            else:
                if not any(v == tf and WANT[kelas] in ctok(l) for l, v in exs) and 'selisih' not in text:
                    err(f'{ur} r{r}: Excel tidak cocok untuk {base!r} {lb!r} {tf} tanpa catatan selisih')
            if per_row:
                nk = ctok(re.sub(r'(?i)^.*?(kelas|-)\s*', '', nm.rsplit(' - ', 1)[-1])) if ' - ' in nm or 'KELAS' in nm.upper() else set()
                if WANT[kelas] not in nk: err(f'{ur} r{r}: kamar {nm!r} bukan kelas {kelas}')
            seq.append(re.sub(r'(?i)\s*-?\s*kelas\s*(III|II|I|1|2|3)\b.*$', '', nm) if per_row else nm)
        seqs.append(seq)
    if not per_row and len(set(map(tuple, seqs))) != 1: err(f'{ur}: urutan/daftar triplet beda antar baris')
    lows = [n.lower() for n in seqs[0] if n]
    if not per_row and len(lows) != len(set(lows)): err(f'{ur}: nama kembar/beda kapital')
print(f'{SHEET}: ERRORS {len(errors)}' + ''.join('\n - ' + e for e in errors[:40]))
sys.exit(1 if errors else 0)
